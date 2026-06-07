from pptx import Presentation
from pptx.util import Inches, Pt

OUT = "c:/Hackathon/hackathon_solution/India_Runs_Presentation.pptx"

slides = [
    ("Title", ["Intelligent Candidate Discovery & Ranking", "Solo – 2026 fresher"], "One-slide summary: heuristic + model ensemble to rank candidates."),
    ("Problem", ["Rank candidates for opportunities using profile and platform signals.", "Produce a top-100 ranked list that meets challenge rules."], "Explain scoring constraints and fairness considerations."),
    ("Data", ["Source: candidates.jsonl", "Signals: completeness, experience, skills, recruiter response, GitHub, views, connections"], "Data fields used and basic cleaning steps."),
    ("Approach", ["Baseline: weighted heuristic", "Model: LightGBM regressor on pseudo-targets", "Ensemble: average baseline + model"], "Why combine heuristic + model? Stability and nonlinear gains."),
    ("Validation", ["Used provided validator: 100 rows, unique ranks, tie-break rules", "Scores rounded to 4 decimals and sorted"], "Show validator checks and how we resolved tie-breaks."),
    ("Results", ["Final file: submission_final.csv", "Demo: top20.json + demo.html"], "Key numbers and proxy metrics on top-100."),
    ("Demo", ["Open demo.html to view top candidates", "Narrate why the top profiles are ranked higher"], "Live demo steps and talking points."),
    ("Why It Works", ["Heuristic captures clear signals", "Model adds nonlinear interactions", "Ensembling stabilizes predictions"], "Short rationale for judges."),
    ("Future", ["Pairwise ranker with feedback", "NLP features from summaries and job history", "More robust hyperparameter tuning"], "Next steps if given more time/data."),
    ("Closing", ["Deliverables: submission_final.csv, submission_final_package.zip, GitHub repo", "Repo: https://github.com/Ayeshasiddiqa789/india-runs-hackathon"], "Contact and how to reproduce quickly."),
]


def add_title_slide(prs, title, subtitle, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle
    if notes:
        ns = s.notes_slide.notes_text_frame
        ns.text = notes


def add_bullets_slide(prs, heading, bullets, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    title = s.shapes.title
    title.text = heading
    # style title
    title.text_frame.paragraphs[0].font.size = Pt(28)
    body = s.shapes.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        p = body.add_paragraph() if i else body.paragraphs[0]
        p.text = b
        p.level = 0
        p.font.size = Pt(18)
    if notes:
        ns = s.notes_slide.notes_text_frame
        ns.text = notes


def build():
    prs = Presentation()
    # Clear default slides
    prs.slides._sldIdLst.clear()
    # Create slides with notes
    add_title_slide(prs, slides[0][1][0], slides[0][1][1], slides[0][2])
    for heading, bullets, notes in slides[1:]:
        add_bullets_slide(prs, heading, bullets, notes)
    prs.save(OUT)
    print(f"Wrote: {OUT}")


if __name__ == '__main__':
    build()
